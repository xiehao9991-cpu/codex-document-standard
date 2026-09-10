"""Regression tests for numeric typography and recursive table coverage."""
import copy
import io
import sys
import unittest
from pathlib import Path
from unittest.mock import patch
from types import SimpleNamespace
import json

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import apply_dingtalk_heading_styles as native


def leaf(text, size=14):
    return ["span", {"data-type": "leaf", "sz": size, "szUnit": "pt", "fonts": {
        "ascii": "Courier New", "hAnsi": "Courier New", "eastAsia": "Microsoft YaHei"},
        "bold": False, "color": "#000000"}, text]


def table(key, fill="#005D8D", rows=None):
    values = rows or [["姓名", "人数"], ["样例", "0012"]]
    return ["table", {"uuid": key}, *[
        ["tr", {}, *[["tc", {"fill": fill if r == 0 else "#FFFFFF"},
                     ["p", {"jc": "left"}, leaf(v)]] for v in row]]
        for r, row in enumerate(values)]]


def apply_offline(root):
    with patch.object(native, "run_dws", return_value='{"success":true}'):
        return native.apply_styles("offline-test", root)


class NativeTests(unittest.TestCase):
    def test_multiline_cell_not_entire_row_or_column_is_left(self):
        t = table("multiline", rows=[["项目", "说明"], ["12", "甲"], ["13", "乙"]])
        cell = t[3][3]
        cell[2].extend([["br", {"type": "textWrapping"}, leaf("")], leaf("第二行")])
        native.normalize_root(["root", {}, t])
        self.assertEqual(cell[2][1]["jc"], "left")
        self.assertEqual(t[3][2][2][1]["jc"], "center")
        self.assertEqual(t[4][3][2][1]["jc"], "center")

    def test_all_paragraphs_in_second_subtable_cell_are_left(self):
        outer = table("outer")
        first, second = table("first"), table("second")
        outer[3][2].extend([first, second, ["p", {}, leaf("")]])
        second[3][3].append(["p", {}, leaf("第二段")])
        native.normalize_root(["root", {}, outer])
        self.assertEqual(outer[3][2][2][1]["jc"], "center")
        self.assertEqual(first[3][3][2][1]["jc"], "center")
        self.assertEqual([p[1]["jc"] for p in second[3][3][2:]], ["left", "left"])

    def test_soft_wrap_uses_each_cells_width(self):
        t = table("wrap", rows=[["窄列", "宽列"], ["这是相同的长文字", "这是相同的长文字"]])
        for row in t[2:]:
            for cell, width in zip(row[2:], (50, 400)):
                cell[1]["width"] = {"type": "dxa", "w": width}
        native.normalize_root(["root", {}, t])
        self.assertEqual(t[3][2][2][1]["jc"], "left")
        self.assertEqual(t[3][3][2][1]["jc"], "center")

    def test_partial_write_is_read_back_before_accepting(self):
        root = ["root", {}, table("t")]
        failure = json.dumps({"error": {"reason": "doc_write_verification_failed", "execution_started": True}})
        with patch.object(native, "run_dws", return_value=failure), patch.object(native, "read_jsonml", side_effect=lambda *a: copy.deepcopy(root)):
            native.apply_styles("offline-test", root)
        root = ["root", {}, table("t")]
        def incorrect_read(*args):
            actual = copy.deepcopy(root)
            next(native.leaf_nodes(actual))[2] = "changed content"
            return actual
        with patch.object(native, "run_dws", return_value=failure), patch.object(native, "read_jsonml", side_effect=incorrect_read):
            with self.assertRaisesRegex(RuntimeError, "partial or changed content"):
                native.apply_styles("offline-test", root)

    def test_cli_verification_error_on_stderr_is_not_a_retry(self):
        error = json.dumps({"error": {"reason": "doc_write_verification_failed", "execution_started": True}})
        result = SimpleNamespace(returncode=1, stdout="", stderr=error)
        with patch.object(native.subprocess, "run", return_value=result) as run:
            self.assertEqual(native.run_dws("doc", "+update", input_text="payload"), error)
            self.assertEqual(run.call_count, 1)
            self.assertEqual(run.call_args.kwargs["input"], "payload")

    def test_visible_merged_cells_centered_and_placeholders_untouched(self):
        t = table("merged", rows=[["分组", "项目", "数量"], ["A", "设计", "12"], ["", "剪辑", "8"]])
        t[3][2][1]["rowSpan"] = 2
        t[4][2][1]["hidden"] = True
        root = ["root", {}, t]
        apply_offline(root)
        self.assertEqual(t[4][4][2][1]["jc"], "center")
        self.assertEqual(t[4][3][2][1]["jc"], "center")
        self.assertNotIn("jc", t[4][2][1])

    def test_color_independent_and_wrapped_tables(self):
        root = ["root", {}, table("t1"), table("t2", "#DCE6F1"),
                ["blockquote", {"uuid": "wrapper"}, table("t3", "#005d8d")]]
        apply_offline(root)
        for t in (root[2], root[3], root[4][2]):
            self.assertTrue(all(n["sz"] == 10 for n in native.iter_leaves(t)))
        native.verify_styles(root, expected_tables=3)

    def test_split_runs_and_alignment(self):
        root = ["root", {}, table("t1")]
        paragraph = root[2][3][3][2]
        paragraph[2:] = [leaf("00"), leaf("12")]
        paragraph[3][1].update({"bold": True, "color": "#FF0000"})
        apply_offline(root)
        self.assertEqual(paragraph[1]["jc"], "center")
        for attrs in native.iter_leaves(paragraph):
            self.assertEqual(attrs["fonts"]["ascii"], "Microsoft YaHei")
            self.assertFalse(attrs["bold"])
            self.assertEqual(attrs["color"], "#000000")
        native.verify_styles(root, expected_tables=1)
        paragraph[3][1]["sz"] = 14
        with self.assertRaises(RuntimeError):
            native.verify_styles(root, expected_tables=1)

    def test_notes_preserved_inside_data_table(self):
        root = ["root", {}, table("outer")]
        note = ["table", {"uuid": "note"}, ["tr", {}, ["tc", {"fill": "#EEF5F8"},
                ["p", {}, leaf("note", 10)]]]]
        nested = table("nested", "#FFFFFF")
        root[2][3][2].extend([note, nested])
        before_note = copy.deepcopy(note)
        apply_offline(root)
        self.assertEqual(note, before_note)
        self.assertTrue(all(n["sz"] == 10 for n in native.iter_leaves(nested)))
        native.verify_styles(root, expected_tables=3)

    def test_no_heading_requirement_and_more_than_thirty_blocks(self):
        root = ["root", {}, *[table(f"t{i}") for i in range(31)]]
        apply_offline(root)
        native.verify_styles(root, expected_tables=31)
        with self.assertRaises(RuntimeError):
            native.verify_styles(root, expected_tables=32)


class LocalTests(unittest.TestCase):
    def test_break_and_multiple_paragraphs_align_whole_cell_left(self):
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH as A
        from normalize_docx_tables import normalize_document
        doc = Document()
        t = doc.add_table(rows=4, cols=2)
        t.cell(1, 1).text = "第一行\n第二行"
        t.cell(2, 1).text = "第一段"
        t.cell(2, 1).add_paragraph("第二段")
        t.cell(3, 1).text = "0012"
        normalize_document(doc)
        self.assertEqual(t.cell(1, 1).paragraphs[0].alignment, A.LEFT)
        self.assertEqual([p.alignment for p in t.cell(2, 1).paragraphs], [A.LEFT, A.LEFT])
        self.assertEqual(t.cell(3, 1).paragraphs[0].alignment, A.CENTER)

    def test_soft_wrap_respects_width_and_cell_margins(self):
        from docx import Document
        from docx.shared import Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH as A
        from normalize_docx_tables import normalize_document
        doc = Document()
        t = doc.add_table(rows=2, cols=2)
        t.autofit = False
        for row in t.rows:
            for cell, width in zip(row.cells, (.5, 3)):
                cell.width = Inches(width)
                cell.text = "这是相同的长文字"
        normalize_document(doc)
        self.assertEqual(t.cell(1, 0).paragraphs[0].alignment, A.LEFT)
        self.assertEqual(t.cell(1, 1).paragraphs[0].alignment, A.CENTER)

    def test_second_subtable_multiline_does_not_left_align_parent_label(self):
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH as A
        from normalize_docx_tables import normalize_document
        doc = Document()
        outer = doc.add_table(rows=2, cols=2)
        parent = outer.cell(1, 0)
        parent.text = "父表"
        for _ in range(2):
            sub = parent.add_table(rows=2, cols=2)
            sub.cell(1, 0).text = "12"
        sub.cell(1, 1).text = "第一行\n第二行"
        normalize_document(doc)
        self.assertEqual(parent.paragraphs[0].alignment, A.CENTER)
        self.assertEqual(sub.cell(1, 1).paragraphs[0].alignment, A.LEFT)
        self.assertEqual(sub.cell(1, 0).paragraphs[0].alignment, A.CENTER)

    def test_table_inside_content_control(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from normalize_docx_tables import normalize_document
        doc = Document()
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "工号", "人数"
        t.cell(1, 0).text, t.cell(1, 1).text = "0012", "12"
        wrapper, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
        t._tbl.addprevious(wrapper)
        wrapper.append(content)
        content.append(t._tbl)
        self.assertEqual(normalize_document(doc)["tables"], 1)

    def test_support_and_merges_preserved_and_each_property_checked(self):
        from docx import Document
        from docx.shared import Pt
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from normalize_docx_tables import normalize_document, verify_document
        doc = Document()
        t = doc.add_table(rows=4, cols=2)
        for ri, values in enumerate((("类别", "数量"), ("分组", ""), ("样例", "12"), ("", "0"))):
            for cell, value in zip(t.rows[ri].cells, values):
                cell.text = value
        t.cell(1, 0).merge(t.cell(1, 1))
        t.cell(2, 0).merge(t.cell(3, 0))
        note = t.cell(2, 0).add_table(rows=1, cols=1)
        note.cell(0, 0).text = "说明123"
        note.cell(0, 0).paragraphs[0].runs[0].font.size = Pt(10)
        note_xml = note._tbl.xml
        merges = list(doc.element.xpath(".//w:gridSpan/@w:val|.//w:vMerge/@w:val"))
        normalize_document(doc)
        self.assertEqual(note._tbl.xml, note_xml)
        self.assertEqual(merges, doc.element.xpath(".//w:gridSpan/@w:val|.//w:vMerge/@w:val"))
        buffer = io.BytesIO()
        doc.save(buffer)
        for prop in ("font", "size", "bold", "color", "spacing", "fill"):
            changed = Document(io.BytesIO(buffer.getvalue()))
            cell = changed.tables[0].cell(2, 1)
            run = cell.paragraphs[0].runs[0]
            if prop == "font":
                run.font.name = "Courier New"
            elif prop == "size":
                run.font.size = Pt(14)
            elif prop == "bold":
                run.font.bold = True
            elif prop == "color":
                from docx.shared import RGBColor
                run.font.color.rgb = RGBColor(255, 0, 0)
            elif prop == "spacing":
                cell.paragraphs[0].paragraph_format.space_after = Pt(12)
            else:
                changed.tables[0].cell(0, 0)._tc.xpath("./w:tcPr/w:shd")[0].set(qn("w:fill"), "FFFFFF")
            with self.subTest(prop=prop), self.assertRaises(RuntimeError):
                verify_document(changed, 2)

    def test_all_data_types_centered(self):
        from table_format_common import column_alignment
        for value in ("0", "-12.50", "1,234.50", "85.50%", "¥1,234.50", "1.2e3"):
            self.assertEqual(column_alignment("数值", [value]), "center")
        for header in ("序号", "日期", "状态"):
            self.assertEqual(column_alignment(header, ["01"]), "center")
        self.assertEqual(column_alignment("工号", ["0012"]), "center")
        self.assertEqual(column_alignment("说明", ["完成12人"]), "center")
        self.assertEqual(column_alignment("空值", ["", "—"]), "center")

    def test_nested_docx_check_and_content_preservation(self):
        from docx import Document
        from docx.shared import Pt
        from normalize_docx_tables import normalize_document, verify_document
        doc = Document()
        doc.add_paragraph("正文123.45")
        outer = doc.add_table(rows=2, cols=2)
        for cell, text in zip(outer.rows[0].cells, ("说明", "人数")):
            cell.text = text
        outer.cell(1, 0).text = "中文12人"
        outer.cell(1, 1).text = "0012"
        for _ in range(2):
            child = outer.cell(1, 0).add_table(rows=2, cols=2)
            child.cell(0, 0).text, child.cell(0, 1).text = "工号", "比例"
            child.cell(1, 0).text, child.cell(1, 1).text = "0012", "85.50%"
        before = [n.text for n in doc.element.xpath(".//w:t")]
        report = normalize_document(doc)
        self.assertEqual(report["tables"], 3)
        stream = io.BytesIO()
        doc.save(stream)
        stream.seek(0)
        saved = Document(stream)
        verify_document(saved, expected_tables=3)
        self.assertEqual(before, [n.text for n in saved.element.xpath(".//w:t")])
        saved.tables[0].cell(1, 0).tables[1].cell(1, 1).paragraphs[0].runs[0].font.size = Pt(14)
        with self.assertRaises(RuntimeError):
            verify_document(saved, expected_tables=3)


if __name__ == "__main__":
    unittest.main()
