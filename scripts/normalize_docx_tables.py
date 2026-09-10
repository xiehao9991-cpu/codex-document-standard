"""Normalize/check document table typography recursively without rebuilding content."""
import argparse
import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.run import Run
from docx.table import Table

from table_format_common import DEFAULT_FONT, LOCAL_SIZE, cell_alignment, row_style, table_role

ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}


def unique_cells(row):
    seen = set()
    for cell in row.cells:
        if cell._tc not in seen:
            seen.add(cell._tc)
            yield cell


def inventory(doc):
    """Enumerate body XML: doc.tables omits nested tables and content controls."""
    found = []
    for index, element in enumerate(doc.element.body.iter(qn("w:tbl")), 1):
        table = Table(element, doc)
        depth = sum(p.tag == qn("w:tbl") for p in element.iterancestors())
        path = f"{index}@depth{depth}"
        cells = [cell for row in table.rows for cell in unique_cells(row)]
        fills = [n.get(qn("w:fill")) for cell in cells for n in cell._tc.xpath("./w:tcPr/w:shd")]
        caption = table._tbl.tblPr.find(qn("w:tblCaption"))
        explicit = caption.get(qn("w:val"), "").removeprefix("codex:") if caption is not None else None
        role = table_role(len(table.rows), len(table.columns), fills,
                          bool(element.xpath(".//w:tbl")), explicit)
        found.append((path, table, role))
    return found


def text_runs(paragraph):
    for element in paragraph._p.xpath(".//w:r"):
        if element.xpath(".//w:t"):
            yield Run(element, paragraph)


def set_font(run, font, size=None, style=None):
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), font)
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{key}"), None)
    if size is not None:
        run.font.size = Pt(size)
        node = rpr.find(qn("w:szCs"))
        if node is None:
            node = OxmlElement("w:szCs")
            rpr.append(node)
        node.set(qn("w:val"), str(int(size * 2)))
    if style is not None:
        run.font.bold = style["bold"]
        run.font.color.rgb = RGBColor.from_string(style["color"])
        color = rpr.find(qn("w:color"))
        for key in ("themeColor", "themeTint", "themeShade"):
            color.attrib.pop(qn(f"w:{key}"), None)


def own_text(cell):
    return "\n".join(p.text for p in cell.paragraphs)


def content_width_pt(cell, table):
    if cell.width is None:
        return None
    margins = []
    for side, alternate in (("left", "start"), ("right", "end")):
        nodes = cell._tc.xpath(f"./w:tcPr/w:tcMar/w:{side}|./w:tcPr/w:tcMar/w:{alternate}")
        if not nodes:
            nodes = table._tbl.xpath(f"./w:tblPr/w:tblCellMar/w:{side}|./w:tblPr/w:tblCellMar/w:{alternate}")
        margins.append(float(nodes[-1].get(qn("w:w"), "108")) / 20 if nodes else 5.4)
    indents = [sum((value.pt if value is not None else 0) for value in
                   (p.paragraph_format.left_indent, p.paragraph_format.right_indent))
               + max(0, p.paragraph_format.first_line_indent.pt if p.paragraph_format.first_line_indent else 0)
               for p in cell.paragraphs]
    return max(0, cell.width.pt - sum(margins) - max(indents, default=0))


def targets(table, font=DEFAULT_FONT):
    if not table.rows:
        return
    headers = [own_text(c) for c in table.rows[0].cells]
    seen = set()
    for ri, row in enumerate(table.rows):
        section = len(list(unique_cells(row))) == 1 and len(headers) > 1
        style = row_style(ri, section)
        for ci, cell in enumerate(row.cells):
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            alignment = cell_alignment([p.text for p in cell.paragraphs], content_width_pt(cell, table),
                                       LOCAL_SIZE, font, style["bold"])
            yield cell, style, alignment


def normal_body_paragraphs(doc):
    for p in doc.paragraphs:
        if p.style.name == "Normal" or p.style.name.startswith("List "):
            yield p


def choose_font(doc, requested=None):
    if requested:
        return requested
    fonts = doc.styles["Normal"]._element.xpath("./w:rPr/w:rFonts")
    return (fonts[0].get(qn("w:eastAsia")) if fonts else None) or DEFAULT_FONT


def normalize_document(doc, font=None):
    font = choose_font(doc, font)
    tables = inventory(doc)
    for p in normal_body_paragraphs(doc):
        for run in text_runs(p):
            set_font(run, font)
    for _, table, role in tables:
        if role != "data":
            continue
        for cell, style, alignment in targets(table, font):
            if style["fill"]:
                tcpr = cell._tc.get_or_add_tcPr()
                shd = tcpr.find(qn("w:shd"))
                if shd is None:
                    shd = OxmlElement("w:shd")
                    tcpr.append(shd)
                shd.set(qn("w:fill"), style["fill"])
            for p in cell.paragraphs:
                p.alignment = ALIGN[alignment]
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for run in text_runs(p):
                    set_font(run, font, LOCAL_SIZE, style)
    return verify_document(doc, expected_tables=len(tables), font=font)


def verify_document(doc, expected_tables=None, font=None):
    font = choose_font(doc, font)
    tables = inventory(doc)
    actual = len(doc.element.body.xpath(".//w:tbl"))
    failures = []
    if len(tables) != actual or (expected_tables is not None and actual != expected_tables):
        failures.append(f"table coverage expected={expected_tables}, traversed={len(tables)}, physical={actual}")
    for path, table, role in tables:
        if role != "data":
            continue
        for cell, style, alignment in targets(table, font):
            if style["fill"]:
                fills = cell._tc.xpath("./w:tcPr/w:shd/@w:fill")
                if not fills or fills[-1].upper() != style["fill"]:
                    failures.append(f"{path}: cell fill")
            for p in cell.paragraphs:
                if p.alignment != ALIGN[alignment]:
                    failures.append(f"{path}: alignment {p.text[:24]}")
                fmt = p.paragraph_format
                if fmt.line_spacing != 1.05 or fmt.space_before != Pt(2) or fmt.space_after != Pt(2):
                    failures.append(f"{path}: paragraph spacing {p.text[:24]}")
                for run in text_runs(p):
                    rpr = run._element.rPr
                    fonts = rpr.rFonts if rpr is not None else None
                    if (run.font.size != Pt(LOCAL_SIZE) or run.font.bold != style["bold"]
                        or str(run.font.color.rgb) != style["color"] or fonts is None
                        or any(fonts.get(qn(f"w:{k}")) != font for k in ("ascii", "hAnsi", "eastAsia", "cs"))
                        or rpr.xpath("./w:szCs/@w:val") != [str(int(LOCAL_SIZE * 2))]):
                        failures.append(f"{path}: text style {run.text[:24]}")
    if failures:
        raise RuntimeError("DOCX verification failed: " + "; ".join(failures[:20]))
    return {"tables": actual, "roles": dict(Counter(role for _, _, role in tables)), "font": font,
            "data_size_pt": LOCAL_SIZE, "soft_wrap_check": "estimated; render review required", "failures": 0}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--check", action="store_true")
    parser.add_argument("--font")
    parser.add_argument("--expected-tables", type=int)
    args = parser.parse_args()
    doc = Document(args.source)
    if args.expected_tables is not None and len(inventory(doc)) != args.expected_tables:
        raise RuntimeError("Source table count differs from --expected-tables")
    if args.check:
        report = verify_document(doc, args.expected_tables, args.font)
    else:
        if args.output is None or args.output.resolve() == args.source.resolve():
            parser.error("Specify a different --output path to preserve the source.")
        original_text = [n.text for n in doc.element.xpath(".//w:t")]
        report = normalize_document(doc, args.font)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        doc.save(args.output)
        saved = Document(args.output)
        if original_text != [n.text for n in saved.element.xpath(".//w:t")]:
            raise RuntimeError("Text content changed")
        report = verify_document(saved, report["tables"], report["font"])
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
